"""Genera Markdown y completa el Word desde el mismo contenido y métricas verificadas.

Ejecutar con el Python del runtime de documentos de Codex. Las figuras son
dibujos científicos de ReportLab rasterizados con PDFium del mismo runtime.
"""
from pathlib import Path
import hashlib
import json
import math
import re
import shutil
import argparse
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.graphics.shapes import Drawing, String, Rect, Line, PolyLine, Polygon
from reportlab.graphics import renderSVG, renderPDF
import pypdfium2 as pdfium
from reportlab.lib import colors

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
WORKSPACE = next(p for p in ROOT.parents if (p / "informe de avance.docx").exists())
QA = ROOT / ".cache/informeAvance"
QA.mkdir(parents=True, exist_ok=True)
DEP = Path(r"C:/Users/aicra/.cache/codex-runtimes/codex-primary-runtime/dependencies")
CIUDADES = [("cali", "Cali"), ("medellin", "Medellín"),
            ("villavicencio", "Villavicencio"), ("ibague", "Ibagué")]
BLUE = colors.HexColor("#17608A")
GRAY = colors.HexColor("#D9D9D9")


def table(headers, rows):
    return "\n".join(["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"]*len(headers)) + " |",
                      *["| " + " | ".join(map(str, row)) + " |" for row in rows]])


def pct(x):
    return "—" if pd.isna(x) else f"{100*float(x):.1f}".replace(".", ",")


def entero(x):
    return f"{int(x):,}".replace(",", ".")


def texto(d, x, y, s, size=9, anchor="start", color=colors.black):
    d.add(String(x, y, str(s), fontName="Helvetica", fontSize=size, textAnchor=anchor, fillColor=color))


def export(d, nombre):
    svg = QA / f"{nombre}.svg"
    renderSVG.drawToFile(d, str(svg))
    dest = BASE / "figuras" / f"{nombre}.png"
    pdf_path = QA / f"{nombre}.pdf"
    renderPDF.drawToFile(d,str(pdf_path))
    pdf = pdfium.PdfDocument(pdf_path)
    pdf[0].render(scale=240/72).to_pil().save(dest)
    pdf.close()
    return dest


def figuras(paneles, runs):
    # Serie de cada municipio, sin unir intervalos ausentes.
    d = Drawing(468, 205)
    for i, (slug, name) in enumerate(CIUDADES):
        panel = paneles[slug]
        ox = (i % 2) * 234
        oy = 104 if i < 2 else 2
        x, y, w, h = ox+36, oy+20, 190, 65
        maximo = math.ceil(panel.casos.max()/100)*100
        texto(d, ox+36, oy+93, name, 10)
        for v in [0, maximo/2, maximo]:
            yy = y + h*v/maximo
            d.add(Line(x, yy, x+w, yy, strokeColor=GRAY, strokeWidth=.4))
            texto(d, x-4, yy-3, int(v), 8, "end")
        for year in [2007, 2013, 2019, 2025]:
            idx = panel.index[panel.anio_epi.eq(year)][0]
            xx = x+w*idx/(len(panel)-1)
            texto(d, xx, y-12, year, 8, "middle")
        points = []
        for j, v in enumerate(panel.casos):
            if pd.isna(v):
                if len(points) >= 4:
                    d.add(PolyLine(points, strokeColor=BLUE, strokeWidth=.65))
                points = []
            else:
                points.extend([x+w*j/(len(panel)-1), y+h*float(v)/maximo])
        if len(points) >= 4:
            d.add(PolyLine(points, strokeColor=BLUE, strokeWidth=.65))
    export(d, "01_series_ciudades")

    d = Drawing(468, 125)
    texto(d, 3, 113, "Implementado", 10)
    nodes = [(3, "Excel", "Fuentes 2007–2025"), (99, "Preparación", "Fechas y semanas"),
             (195, "Modelos", "Etiqueta y evaluación"), (291, "MLflow en Docker", "Métricas y artefactos")]
    for x, t1, t2 in nodes:
        d.add(Rect(x, 64, 88, 38, fillColor=colors.HexColor("#EAF2F8"), strokeColor=BLUE, strokeWidth=.7))
        texto(d, x+44, 86, t1, 9, "middle")
        texto(d, x+44, 72, t2, 7.8, "middle")
        if x < 291:
            d.add(Line(x+89, 83, x+94, 83, strokeColor=BLUE))
            d.add(Polygon([x+95,83,x+91,85,x+91,81], fillColor=BLUE, strokeColor=BLUE))
    texto(d, 3, 45, "Pendiente", 10)
    for x, label in [(99,"API Python"),(291,"Interfaz React")]:
        d.add(Rect(x, 4, 88, 30, fillColor=colors.white, strokeColor=colors.grey, strokeDashArray=[3,2]))
        texto(d, x+44, 16, label, 9, "middle")
    d.add(Line(335,64,335,44,strokeColor=colors.grey,strokeDashArray=[3,2]))
    d.add(Line(335,44,143,44,strokeColor=colors.grey,strokeDashArray=[3,2]))
    d.add(Line(143,44,143,34,strokeColor=colors.grey,strokeDashArray=[3,2]))
    d.add(Line(291,19,187,19,strokeColor=colors.grey,strokeDashArray=[3,2]))
    texto(d, 239, 23, "consulta y respuesta", 8, "middle")
    texto(d, 393, 82, "SQLite", 9)
    texto(d, 393, 69, "Volumen local", 8)
    export(d, "02_arquitectura")

    d = Drawing(468, 225)
    methods = [("logistica_umbral_validado","Log validada"),("logistica_umbral_0_5","Log 0,5"),
               ("persistencia_conteo","Persistencia"),("prevalencia_entrenamiento","Prev constante"),
               ("siempre_alerta","Siempre alerta")]
    palette = ["#17608A", "#6D9FBD", "#34805E", "#ACB5BD", "#C79436"]
    for i, (_, name) in enumerate(CIUDADES):
        ox, oy = (i % 2)*234, 113 if i < 2 else 0
        texto(d, ox+4, oy+103, name, 10)
        for v in [0, 50, 100]:
            xx = ox+78+1.25*v
            d.add(Line(xx,oy+22,xx,oy+94,strokeColor=GRAY,strokeWidth=.45))
            texto(d, xx, oy+11, v, 8, "middle")
        for k, ((method,label), color) in enumerate(zip(methods,palette)):
            value = runs[(name,method)]["metrics"]["validacion.f2"]*100
            yy=oy+82-k*14
            texto(d, ox+73, yy+1, label, 8.2, "end")
            d.add(Rect(ox+78, yy-1, value*1.25, 9, fillColor=colors.HexColor(color), strokeColor=None))
            texto(d, ox+82+value*1.25, yy, f"{value:.1f}".replace(".",","), 8)
    export(d, "03_f2_validacion_mlflow")

    d = Drawing(468, 225)
    for i, (_, name) in enumerate(CIUDADES):
        m=runs[(name,"logistica_umbral_validado")]["metrics"]
        vals=[[m['prueba_2024.tn'],m['prueba_2024.fp']],[m['prueba_2024.fn'],m['prueba_2024.tp']]]
        ox,oy=(i % 2)*234,113 if i<2 else 0
        texto(d,ox+5,oy+103,f"{name}   n={int(m['prueba_2024.n'])}",10)
        for row in range(2):
            texto(d,ox+67,oy+72-row*29,["Sin exceso","Exceso"][row],8.5,"end")
            for col in range(2):
                value=vals[row][col]
                t=value/max(v for vv in vals for v in vv)
                d.add(Rect(ox+74+col*55,oy+59-row*29,55,29,
                    fillColor=colors.Color(1-.65*t,1-.42*t,1-.20*t),strokeColor=colors.HexColor('#CBD5DF'),strokeWidth=.5))
                texto(d,ox+101.5+col*55,oy+69-row*29,int(value),12,"middle")
        for col,label in enumerate(["Sin alerta","Alerta"]):
            texto(d,ox+101.5+col*55,oy+17,label,8.5,"middle")
    export(d,"04_matrices_mlflow_2024")


def inline(paragraph, text):
    pattern=r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))"
    for part in re.split(pattern,text):
        if not part:
            continue
        link=re.fullmatch(r"\[([^\]]+)\]\(([^\)]+)\)",part)
        if link:
            rel=paragraph.part.relate_to(link.group(2),RT.HYPERLINK,is_external=True)
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'17608A');pr.append(c);r.append(pr)
            t=OxmlElement('w:t');t.text=link.group(1);r.append(t);h.append(r);paragraph._p.append(h)
        else:
            run=paragraph.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
            if part.startswith('**'):run.bold=True
            if part.startswith('`'):run.font.name='Consolas';run.font.size=Pt(9.5)


def format_table(t, widths=None):
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    n=len(t.columns)
    if widths is None:
        widths=([3.5,3.0] if n==2 else [6.5/n]*n)
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge,val in [('top',55),('bottom',55),('left',75),('right',75)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8F1' if i==0 else 'F5F8FA' if i%2==0 else 'FFFFFF');pr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if n==2 or j==0 or (n==7 and j==1) else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size=Pt(9);run.bold=i==0;run.font.color.rgb=RGBColor(0,0,0)


def plaintext(text):
    text=re.sub(r"\[([^\]]+)\]\([^\)]+\)",r"\1",text)
    return text.replace('**','').replace('`','')


def word_from_md(md, target):
    backup=QA/'borrador_original.docx'
    if not backup.exists():shutil.copy2(target,backup)
    doc=Document(backup)
    for el in list(doc._element.body):
        if el.tag!=qn('w:sectPr'):doc._element.body.remove(el)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    normal.font.color.rgb=RGBColor(0,0,0)
    pf=normal.paragraph_format;pf.space_after=Pt(6);pf.line_spacing=1.02
    for name,size in [('Title',16),('Heading 1',12.5),('Heading 2',11.5),('Caption',9)]:
        if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(7 if name.startswith('Heading') else 0)
        st.paragraph_format.space_after=Pt(5)
        st.paragraph_format.keep_with_next=name.startswith('Heading')
        if name.startswith('Heading') or name=='Title':st.font.bold=True
    doc.styles['Title'].paragraph_format.keep_with_next=True
    footer=doc.sections[0].footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    lines=md.splitlines();i=0;expected=[];salto_pendiente=False
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='<!-- pagina -->':salto_pendiente=True;i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells):rows.append(cells)
                i+=1
            t=doc.add_table(rows=len(rows),cols=len(rows[0]))
            for r,vals in zip(t.rows,rows):
                for c,txt in zip(r.cells,vals):inline(c.paragraphs[0],txt);expected.append(plaintext(txt))
            header=rows[0][0]
            widths=None
            if header=='Ciudad' and len(rows[0])==7:widths=[1.15,1.3,.85,.85,.85,.75,.75]
            if header=='Ciudad' and len(rows[0])==6:widths=[1.1,1.1,1.1,1.0,1.1,1.1]
            if header=='Ciudad' and len(rows[0])==5:widths=[1.3,1.3,1.3,1.3,1.3]
            if header=='Fecha':widths=[1.2,4.4,.9]
            format_table(t,widths)
            continue
        image=re.fullmatch(r'!\[([^\]]+)\]\(([^\)]+)\)',line)
        if image:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
            pic=p.add_run().add_picture(str(BASE/image.group(2)),width=Inches(6.5))
            pic._inline.docPr.set('descr',image.group(1));i+=1;continue
        level=len(line)-len(line.lstrip('#'))
        content=line[level:].strip() if level else line
        style='Title' if level==1 else 'Heading 1' if level==2 else 'Heading 2' if level==3 else 'Caption' if line.startswith(('Figura ','Tabla ')) else 'Normal'
        p=doc.add_paragraph(style=style)
        if salto_pendiente:
            p.paragraph_format.page_break_before=True
            salto_pendiente=False
        inline(p,content);expected.append(plaintext(content))
        if line.startswith('[') and re.match(r'^\[\d+\]',line):
            p.paragraph_format.space_after=Pt(5)
            for r in p.runs:r.font.size=Pt(9.5)
        i+=1
    doc.core_properties.title='Sistemas de alertas tempranas para detectar excesos de dengue en Colombia'
    doc.core_properties.subject='Informe de avance al 24 de octubre de 2026'
    doc.core_properties.author='Fabian Andrés Gullaván; Sebastian Caballero; Jhon Alejandro Zuluaga; Juan Sebastian Garcia'
    doc.save(target)
    # Igualdad semántica entre las representaciones, incluidos párrafos de tablas e hipervínculos.
    from zipfile import ZipFile
    from lxml import etree
    with ZipFile(target) as z:xml=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    actual=[''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in xml.xpath('.//w:body//w:p',namespaces=ns)]
    actual=[v for v in actual if v]
    expected=[v for v in expected if v]
    assert actual==expected,'Diferencias de texto entre Markdown y Word'
    return len(expected)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--salida-word', type=Path, help='Ruta de salida; por defecto informeAvance/informe_avance.docx')
    args=parser.parse_args()
    data=json.loads((BASE/'datos/ejecuciones_mlflow.json').read_text(encoding='utf-8-sig'))
    assert len(data['runs'])==20 and all(r['status']=='FINISHED' for r in data['runs'])
    runs={(r['tags']['ciudad'],r['tags']['metodo']):r for r in data['runs']}
    assert len(runs)==20
    perfiles=[];paneles={};splits=[];rows=[]
    for slug,name in CIUDADES:
        output=ROOT if slug=='cali' else ROOT/'ciudades'/slug
        q=json.loads((output/'reports/resumen_eda.json').read_text(encoding='utf-8'))
        perfiles.append({'ciudad':name,**{k:q[k] for k in ['filas_extraidas','filas_finales','ini_sin_no_parseable','semanas_sin_registros_no_imputadas','semanas_elegibles_modelo']}})
        paneles[slug]=pd.read_csv(output/f'data/processed/{slug}_panel_semanal.csv')
        s=pd.read_csv(output/'reports/tables/particiones_temporales.csv');s.insert(0,'ciudad',name);splits.append(s)
        local=pd.read_csv(output/'reports/tables/metricas_logistica.csv')
        for r in [v for (city,_),v in runs.items() if city==name]:
            method=r['tags']['metodo']
            for v in local.loc[local.modelo.eq(method)].to_dict('records'):
                for k,val in v.items():
                    if k in ['modelo','bloque']:continue
                    key=f"{v['bloque']}.{k}"
                    if pd.isna(val):assert key not in r['metrics']
                    else:assert math.isclose(r['metrics'][key],float(val),abs_tol=1e-12,rel_tol=0),(name,method,key)
                rows.append({'ciudad':name,'run_id':r['run_id'],**v})
    pd.DataFrame(rows).to_csv(BASE/'datos/metricas_mlflow.csv',index=False)
    pd.DataFrame(perfiles).to_csv(BASE/'datos/perfil_ciudades.csv',index=False)
    s=pd.concat(splits,ignore_index=True);s.to_csv(BASE/'datos/particiones_temporales.csv',index=False)
    figuras(paneles,runs)
    replacements={}
    replacements['tabla_calidad']=table(['Ciudad','Extraídos','Incluidos','Sem sin filas','Sem elegibles'],[
        [p['ciudad'],entero(p['filas_extraidas']),entero(p['filas_finales']),p['semanas_sin_registros_no_imputadas'],p['semanas_elegibles_modelo']] for p in perfiles])
    replacements['tabla_variables']=table(['Grupo de variables','Construcción y disponibilidad'],[
        ['Rezagos 1 2 3 4 8 y 12','Seis conteos de semanas anteriores al objetivo.'],
        ['Media y desviación de cuatro semanas','Dos resúmenes calculados después de desplazar una semana.'],
        ['Seno y coseno de la semana','Dos variables del ciclo anual conocido, considerando años de 52 o 53 semanas.']])
    part=[]
    for _,name in CIUDADES:
        b=s.loc[s.ciudad.eq(name)].set_index('bloque')
        part.append([name,*[f"{int(b.loc[k,'filas'])} ({int(b.loc[k,'positivos'])})" for k in ['entrenamiento','validacion','prueba_2024','adicional_2025']]])
    replacements['tabla_particiones']=table(['Ciudad','Entrenamiento 2012–2021','Validación 2022–2023','Prueba 2024','Adicional 2025'],part)
    val=[]
    for _,name in CIUDADES:
        for method,label in [('logistica_umbral_validado','Log seleccionada'),('persistencia_conteo','Persistencia')]:
            m=runs[(name,method)]['metrics'];val.append([name,label,*[pct(m.get('validacion.'+k,float('nan'))) for k in ['sensibilidad','precision','f1']],int(m['validacion.fp']),int(m['validacion.fn'])])
    replacements['tabla_validacion']=table(['Ciudad','Método','Sens %','Prec %','F1 %','FP','FN'],val)
    posterior=[]
    for _,name in CIUDADES:
        log=runs[(name,'logistica_umbral_validado')]['metrics'];pers=runs[(name,'persistencia_conteo')]['metrics']
        posterior.append([name,pct(log.get('prueba_2024.f1',float('nan'))),pct(pers.get('prueba_2024.f1',float('nan'))),pct(log.get('adicional_2025.f1',float('nan'))),pct(pers.get('adicional_2025.f1',float('nan'))),int(log['adicional_2025.fp'])])
    replacements['tabla_posterior']=table(['Ciudad','F1 Log 2024','F1 Pers 2024','F1 Log 2025','F1 Pers 2025','FP Log 2025'],posterior)
    replacements['tabla_cronograma']=table(['Fecha','Hito','Estado'],[
        ['25 oct','Entregar avance con evidencia y registro de experimentos','En curso'],
        ['26 oct–8 nov','Optimizar logística y comparar HistGradientBoosting y MLP','Previsto'],
        ['9–19 nov','Implementar API Python e interfaz React y sus pruebas','Previsto'],
        ['20–26 nov','Integrar, validar y desplegar la solución el 26 de noviembre','Previsto'],
        ['27 nov','Sustentación del proyecto','Previsto'],
        ['29 nov','Entrega del informe final','Previsto']])
    md=(BASE/'contenido_base.md').read_text(encoding='utf-8')
    for key,value in replacements.items():md=md.replace('{{'+key+'}}',value)
    assert '{{' not in md
    resumen=md.split('### Resumen\n\n')[1].split('\n\nPalabras clave:')[0]
    words=len(resumen.split());assert 200<=words<=300,words
    assert md.count('<!-- pagina -->')==7
    (BASE/'informe_avance.md').write_text(md,encoding='utf-8')
    target=args.salida_word.resolve() if args.salida_word else BASE/'informe_avance.docx'
    paragraphs=word_from_md(md,target)
    status={'fecha_corte':'2026-10-24','resumen_palabras':words,'bloques_planificados':8,
        'parrafos_y_celdas_cotejados':paragraphs,'ejecuciones_mlflow_verificadas':20,
        'metricas_y_csv_locales':'coinciden','versiones_word_markdown':'mismo contenido',
        'sha256_snapshot_mlflow':hashlib.sha256((BASE/'datos/ejecuciones_mlflow.json').read_bytes()).hexdigest(),
        'word':str(target),'sha256_word':hashlib.sha256(target.read_bytes()).hexdigest()}
    (QA/'verificacion_contenido.json').write_text(json.dumps(status,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(status,ensure_ascii=True,indent=2))


if __name__=='__main__':main()
